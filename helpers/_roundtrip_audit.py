#!/usr/bin/env python3
"""
Scratch audit (2026-09-13): every session folder under the Ext roundtrip image
tree -> one status, cross-checked against the dynamic recon outputs.

Status (first match wins):
  DONE          output folder in AIkill_Dynamic with >=1 method output.mrd
  QUEUED        in tonight's resumed batch, not finished yet
  FAILED        latest SUMMARY.txt entry rc!=0 and no outputs (reason from its log)
  DUPLICATE     largest dynamic .dat identical (size + md5 of first 20 MB) to a DONE/other folder
  NO_TRAJ       spiral-dyn .dat present but no gp trajectory for its seqname
  NOT_DYNAMIC   no spiral-dyn .dat at all
  NOT_RUN       has a runnable dynamic .dat but none of the above (should be empty)

Writes CSV + summary to workspace/outputs/roundtrip_audit_2026-09-13/.
"""
import csv
import hashlib
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

IMAGES = Path('/Volumes/HoomHamExt/_5t_images_roundtrip/Images')
DYN_OUT = Path('/Volumes/HoomHamExt/AIkill_Dynamic')
LOGS = Path('/Volumes/HoomHamExt/Work/Codes/2026_ASAP_Recon/pipeline_runs/logs')
WS = Path(__file__).resolve().parents[1]
TRAJ_DIR = WS / 'data' / 'xe' / 'human' / 'traj'
SNR_XLSX = WS / 'data' / 'SNR_Table_All.xlsx'
OUT_DIR = WS / 'outputs' / 'roundtrip_audit_2026-09-13'
RUNNER_OUT = LOGS / 'runner_resume.out'

SEQ_RE = re.compile(r'fa_spiral_dyn_fancy_(v\d_\d{8}|\d{8})')
DATE_RE = re.compile(r'(\d{4}-\d{2}-\d{2})')
# staged loose-layout sessions reconstructed from other Ext copies
STAGED = {'2023-07-24_019WR', '2023-11-02_000LL'}


def real_files(d):
    try:
        return sorted(f for f in os.listdir(d) if not f.startswith('._') and not f.startswith('.'))
    except (PermissionError, FileNotFoundError, NotADirectoryError):
        return []


def md5_head(p, nbytes=20_000_000):
    h = hashlib.md5()
    with open(p, 'rb') as f:
        h.update(f.read(nbytes))
    return h.hexdigest()


def traj_available(seq):
    return (TRAJ_DIR / f'fa_spiral_dyn_fancy_{seq}_gp.npy').is_file()


# ---- SNR table rows ----
snr_rows = defaultdict(list)
wb = openpyxl.load_workbook(SNR_XLSX)
ws = wb['Sheet1']
for r in range(3, ws.max_row + 1):
    dt, sid = ws.cell(r, 2).value, ws.cell(r, 3).value
    if not (dt and sid):
        continue
    if hasattr(dt, 'strftime'):
        dt = dt.strftime('%Y-%m-%d')
    else:
        y, m, d = re.split(r'[-/]', str(dt))
        dt = f'{int(y):04d}-{int(m):02d}-{int(d):02d}'
    snr_rows[f'{dt}_{sid}'].append(r)

# ---- latest batch summary per session ----
summary = {}
if (LOGS / 'SUMMARY.txt').is_file():
    for line in (LOGS / 'SUMMARY.txt').read_text().splitlines():
        parts = line.split()
        if len(parts) >= 4:
            summary[parts[1]] = (parts[2], parts[3])   # rc=, outputs=

# ---- sessions still queued in the resumed runner ----
queued = set()
if RUNNER_OUT.is_file():
    txt = RUNNER_OUT.read_text()
    started = set(re.findall(r'START\s+(\S+)', txt))
    ended = set(re.findall(r'END\s+(\S+)', txt))
    in_progress = started - ended
    worklist = ['2023-03-09_023DB', '2023-03-10_004LR', '2023-03-24_029CK',
                '2023-03-27_02BB', '2023-03-31_032WS']
    queued = (set(worklist) - ended) | in_progress


def fail_reason(key):
    log = LOGS / f'{key}.log'
    if not log.is_file():
        return ''
    lines = [l for l in log.read_text(errors='replace').splitlines()
             if re.search(r'Error|ERROR', l) and 'read data' not in l]
    return lines[-1].strip()[-120:] if lines else ''


def output_status(key):
    d = DYN_OUT / key
    if not d.is_dir():
        return '', ''
    methods = ''.join(m if (d / m / 'output.mrd').is_file() else '-' for m in 'spd')
    kinds = set()
    for m in 'spd':
        log = d / m / 'tyger.log'
        if not log.is_file():
            continue
        t = log.read_text(errors='replace')
        if 'identified sampling pattern gas-only' in t:
            kinds.add('gas-only-acq')
        elif 'unsplit complex' in t:
            kinds.add('dissolved-unsplit')
        elif 'RBC/TP phase solved' in t:
            kinds.add('dissolved-split')
        elif 'skipping DPDYN' in t:
            kinds.add('dissolved-SKIPPED(old guard)')
    return methods, ','.join(sorted(kinds))


# ---- walk ----
rows = []
for top in sorted(p for p in IMAGES.iterdir() if p.is_dir() and not p.name.startswith('.')):
    tfiles = real_files(top)
    loose_dat = [f for f in tfiles if f.endswith('.dat') and (top / f).is_file()]
    subdirs = [top / f for f in tfiles if (top / f).is_dir()]
    date_m = DATE_RE.search(top.name)
    date = date_m.group(1) if date_m else top.name
    layout = 'ok' if date_m and top.name == date else f'odd top folder "{top.name}"'
    if loose_dat:
        subdirs = [top] + subdirs
        layout = f'{len(loose_dat)} loose .dat at date level'
    if not subdirs:
        rows.append(dict(folder=top.name, key='', status='EMPTY', note='no session subfolder'))
        continue
    for sd in subdirs:
        sid = sd.name if sd != top else '(date-level)'
        files = real_files(sd)
        dats = [f for f in files if f.endswith('.dat') and (sd / f).is_file()]
        dyn = [f for f in dats if 'spiral_dyn' in f]
        seqs = sorted({SEQ_RE.search(f).group(1) for f in dyn if SEQ_RE.search(f)})
        pneumo = any('pneumotach' in f.lower() and not f.endswith('.dat') for f in files)
        key = f'{date}_{sid}'
        big = max(dyn, key=lambda f: (sd / f).stat().st_size) if dyn else None
        big_mb = round((sd / big).stat().st_size / 1e6) if big else 0
        methods, kinds = output_status(key)
        rc, outs = summary.get(key, ('', ''))
        row = dict(folder=str(sd.relative_to(IMAGES)), key=key, layout=layout,
                   n_dat=len(dats), n_dyn=len(dyn), seqs='+'.join(seqs), dyn_mb=big_mb,
                   pneumotach='yes' if pneumo else 'no',
                   snr_rows=','.join(map(str, snr_rows.get(key, []))),
                   outputs=methods, dissolved=kinds, last_batch=f'{rc} {outs}'.strip(),
                   big_dat=big or '', md5='', status='', note='')
        if big:
            row['md5'] = f'{(sd / big).stat().st_size}:{md5_head(sd / big)}'
        rows.append(row)

# duplicates: same size+md5 in more than one folder
by_md5 = defaultdict(list)
for r in rows:
    if r.get('md5'):
        by_md5[r['md5']].append(r)

for r in rows:
    if r.get('status') == 'EMPTY':
        continue
    key = r['key']
    dups = [o for o in by_md5.get(r.get('md5'), []) if o is not r]
    has_traj = any(traj_available(s) for s in r['seqs'].split('+') if s)
    if r['outputs'] and r['outputs'] != '---':
        r['status'] = 'DONE'
        if dups:
            r['note'] = 'also identical to ' + ', '.join(o['folder'] for o in dups)
    elif key in queued:
        r['status'] = 'QUEUED'
    elif r['last_batch'].startswith('rc=') and not r['last_batch'].startswith('rc=0'):
        r['status'] = 'FAILED'
        r['note'] = fail_reason(key)
    elif dups:
        r['status'] = 'DUPLICATE'
        r['note'] = 'identical dyn .dat to ' + ', '.join(o['folder'] for o in dups)
    elif not r['n_dyn']:
        r['status'] = 'NOT_DYNAMIC'
    elif not has_traj:
        r['status'] = 'NO_TRAJ'
        r['note'] = f'no gp trajectory for {r["seqs"]}'
    else:
        r['status'] = 'NOT_RUN'

OUT_DIR.mkdir(parents=True, exist_ok=True)
fields = ['status', 'folder', 'key', 'layout', 'n_dat', 'n_dyn', 'seqs', 'dyn_mb', 'pneumotach',
          'snr_rows', 'outputs', 'dissolved', 'last_batch', 'note', 'big_dat']
with open(OUT_DIR / 'roundtrip_audit.csv', 'w', newline='') as f:
    w = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore')
    w.writeheader()
    for r in sorted(rows, key=lambda r: (r['status'], r['folder'])):
        w.writerow(r)

# outputs that don't map to any roundtrip folder (e.g. staged from other copies)
keys = {r['key'] for r in rows}
orphans = sorted(p.name for p in DYN_OUT.iterdir() if p.is_dir() and p.name not in keys)

cnt = Counter(r['status'] for r in rows)
print(f'folders audited: {len(rows)}')
for s, n in sorted(cnt.items()):
    print(f'  {s:12s} {n}')
print('\nDissolved type among DONE:')
print('  ' + str(Counter(r['dissolved'] for r in rows if r['status'] == 'DONE')))
for s in ('QUEUED', 'FAILED', 'DUPLICATE', 'NO_TRAJ', 'NOT_DYNAMIC', 'NOT_RUN', 'EMPTY'):
    sel = [r for r in rows if r['status'] == s]
    if not sel:
        continue
    print(f'\n== {s} ({len(sel)})')
    for r in sorted(sel, key=lambda r: r['folder']):
        print(f"  {r['folder']:28s} seq={r.get('seqs','') or '-':22s} dyn={r.get('dyn_mb',0)}MB "
              f"snr_rows={r.get('snr_rows','') or '-'}  {r.get('note','')[:90]}")
odd = [r for r in rows if r.get('layout', 'ok') != 'ok']
if odd:
    print('\n== odd layout')
    for r in odd:
        print(f"  {r['folder']:28s} {r['layout']}  -> {r['status']}")
partial = [r for r in rows if r['status'] == 'DONE' and '-' in r['outputs']]
print(f'\n== DONE but not all 3 methods ({len(partial)})')
for r in partial:
    print(f"  {r['key']:22s} outputs={r['outputs']} pneumotach={r['pneumotach']}")
print(f'\n== outputs with no matching roundtrip folder: {orphans}')
print(f'\nCSV -> {OUT_DIR / "roundtrip_audit.csv"}')
