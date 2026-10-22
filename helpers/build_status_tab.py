#!/usr/bin/env python3
"""Build a 'Recon Status' tab in SNR_Table_All.xlsx.

Classifies every ungreened subject by why it is not reconstructed, and links
each entry back to its source row on Sheet1 via an internal HYPERLINK.
"""
import os, re, datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

XLSX = Path('/Users/hoomham/Hooman/Work/Codes/2026_ASAP_Recon/workspace/data/SNR_Table_All.xlsx')
CUR  = Path('/Volumes/HoomHam/Images/human')            # live source drive
ARC  = Path('/Volumes/HooMain/LeBigMac_HD4/Images/MRI/Human')  # archive drive
OUT  = Path('/Volumes/HoomHamExt/Dynamic')              # recon outputs
TRAJ = Path('/Users/hoomham/Hooman/Work/Codes/2026_ASAP_Recon/workspace/data/xe/human/traj')

KNOWN = {'fa_spiral_dyn_fancy_v2_20230131',
         'fa_spiral_dyn_fancy_v3_20230821',
         'fa_spiral_dyn_fancy_v3_20240130'}
CONTAINER_FAIL = {17: 'curve_fit non-convergence', 46: 'dissolved-phase spectral fit'}

SEQ_RE = re.compile(r'(fa_spiral_dyn_fancy(?:_v\d)?_\d{8})')


def norm_date(d):
    if hasattr(d, 'strftime'):
        return d.strftime('%Y-%m-%d')
    p = str(d).strip().replace('/', '-').split('-')
    return f'{int(p[0]):04d}-{int(p[1]):02d}-{int(p[2]):02d}' if len(p) == 3 else str(d)


def id_key(s):
    """'001BB' -> ('BB', 1); '02BB' -> ('BB', 2). Letter suffix is the stable part."""
    m = re.match(r'^(\d*)([A-Za-z]+)$', str(s).strip())
    return (m.group(2).upper(), int(m.group(1) or -1)) if m else (str(s).upper(), -1)


def listdir(p):
    try:
        return [f for f in os.listdir(p) if not f.startswith('._')]
    except OSError:
        return []


def probe(folder):
    """Return (seqs, has_pneumotach, n_dats) for a raw subject folder."""
    fs = listdir(folder)
    dats = [f for f in fs if f.endswith('.dat')]
    seqs = set()
    for f in dats:
        if 'spiral_dyn' in f.lower():
            m = SEQ_RE.search(f)
            seqs.add(m.group(1) if m else '?')
    pneu = any('pneumotach' in f.lower() and not f.endswith('.dat') for f in fs)
    return seqs, pneu, len(dats)


def find_data(date, sid):
    """Locate raw data. Returns (drive, folder_path, disk_id) or (None, None, None)."""
    # 1. live drive, exact
    if (CUR / date / sid).is_dir():
        return 'HoomHam', CUR / date / sid, sid
    # 2. live drive, fuzzy id (letter-suffix match)
    dd = CUR / date
    if dd.is_dir():
        want = id_key(sid)[0]
        for x in listdir(dd):
            if (dd / x).is_dir() and id_key(x)[0] == want:
                return 'HoomHam', dd / x, x
    # 3. archive drive, <date>_<id> or <date>-<id>
    for name in listdir(ARC):
        if date in name and sid in name:
            return 'HooMain', ARC / name, sid
    # 4. archive drive, fuzzy id
    want = id_key(sid)[0]
    for name in listdir(ARC):
        if date in name:
            tail = name.replace(date, '').lstrip('_-')
            if tail and id_key(tail)[0] == want:
                return 'HooMain', ARC / name, tail
    return None, None, None


def has_output(date, disk_id):
    if not disk_id:
        return False
    base = OUT / f'{date}_{disk_id}'
    return any((base / m / 'output.mrd').is_file() or (base / m / 'recon.mat').is_file() for m in 'spd')


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Sheet1']

    def is_green(r):
        f = ws.cell(r, 2).fill
        return f.fill_type == 'solid' and (f.start_color.rgb or '') == 'FF92D050'

    recs = []
    n_green = 0
    for r in range(3, 149):
        d, sid = ws.cell(r, 2).value, ws.cell(r, 3).value
        if not d or not sid:
            continue
        date = norm_date(d)
        snr = ws.cell(r, 5).value
        if is_green(r):
            n_green += 1
            continue

        drive, folder, disk_id = find_data(date, sid)
        seqs, pneu, ndat = (probe(folder) if folder else (set(), False, 0))
        seq_s = ', '.join(sorted(seqs)) if seqs else ''
        runnable_seq = seqs & KNOWN

        # --- classify -------------------------------------------------
        if drive and has_output(date, disk_id):
            status, blocker = 'RECONSTRUCTED', 'Output exists; row not yet green-marked'
        elif r in CONTAINER_FAIL:
            status, blocker = 'CONTAINER FAIL', f'{CONTAINER_FAIL[r]} — untested vs d136eb1 guard'
        elif runnable_seq:
            status = 'RUNNABLE'
            blocker = 'numba regression in d136eb1 image — pin deps, rebuild'
            if drive == 'HooMain':
                blocker = 'Copy to HoomHam first, then ' + blocker
        elif seqs:
            status, blocker = 'NO TRAJECTORY', f'No {sorted(seqs)[0]}_gp.npy — request from Faraz'
        elif drive:
            status, blocker = 'NO DYNAMIC SCAN', f'{ndat} .dat, none spiral-dyn — never acquired'
        else:
            status, blocker = 'NO DATA', 'Not found on HoomHam or HooMain'

        traj = 'yes' if runnable_seq else ('no' if seqs else '')
        recs.append(dict(row=r, date=date, sid=sid, disk=disk_id or '', snr=snr,
                         status=status, drive=drive or '', seq=seq_s, traj=traj,
                         pneu=('yes' if pneu else ('no' if drive else '')),
                         blocker=blocker))

    # --- order: most actionable first ---------------------------------
    ORDER = ['RECONSTRUCTED', 'RUNNABLE', 'CONTAINER FAIL',
             'NO TRAJECTORY', 'NO DYNAMIC SCAN', 'NO DATA']
    recs.sort(key=lambda x: (ORDER.index(x['status']),
                             -(x['snr'] if isinstance(x['snr'], (int, float)) else 0)))

    # --- styles --------------------------------------------------------
    FILL = {
        'RECONSTRUCTED':   'FF92D050',
        'RUNNABLE':        'FFFFD966',
        'CONTAINER FAIL':  'FFF4B183',
        'NO TRAJECTORY':   'FFFF9999',
        'NO DYNAMIC SCAN': 'FFD9D9D9',
        'NO DATA':         'FFBFBFBF',
    }
    hdr_fill = PatternFill('solid', start_color='FF404040')
    hdr_font = Font(bold=True, color='FFFFFFFF')
    thin = Side(style='thin', color='FFB0B0B0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    name = 'Recon Status'
    if name in wb.sheetnames:
        del wb[name]
    st = wb.create_sheet(name, 1)

    # --- summary block --------------------------------------------------
    counts = {s: sum(1 for x in recs if x['status'] == s) for s in ORDER}
    st['A1'] = 'ASAP Recon — Subject Status'
    st['A1'].font = Font(bold=True, size=14)
    st['A2'] = f'Generated {datetime.date.today().isoformat()}  ·  Sheet1 rows 3–148'
    st['A2'].font = Font(italic=True, color='FF808080')

    st['A4'] = 'Status'; st['B4'] = 'Count'; st['C4'] = 'Meaning'
    for c in 'ABC':
        st[f'{c}4'].font = hdr_font
        st[f'{c}4'].fill = hdr_fill
    MEANING = {
        'RECONSTRUCTED':   'Output on HoomHamExt; needs green-marking on Sheet1',
        'RUNNABLE':        'Data + trajectory present. Blocked only by the numba image',
        'CONTAINER FAIL':  'Crashed in cloud recon; not retried against the guard',
        'NO TRAJECTORY':   'Dynamic scan exists, no matching _gp.npy',
        'NO DYNAMIC SCAN': 'Subject has .dat files but no spiral-dyn acquisition — unrecoverable',
        'NO DATA':         'No raw folder on either drive',
    }
    row = 5
    st.cell(row, 1, 'DONE (green)').fill = PatternFill('solid', start_color='FF92D050')
    st.cell(row, 2, n_green)
    st.cell(row, 3, 'Already reconstructed and green-marked on Sheet1')
    row += 1
    for s in ORDER:
        st.cell(row, 1, s).fill = PatternFill('solid', start_color=FILL[s])
        st.cell(row, 2, counts[s])
        st.cell(row, 3, MEANING[s])
        row += 1
    st.cell(row, 1, 'TOTAL').font = Font(bold=True)
    st.cell(row, 2, n_green + len(recs)).font = Font(bold=True)

    # --- detail table ----------------------------------------------------
    top = row + 2
    HEADERS = ['Sheet1 Row', 'Date', 'ID (sheet)', 'ID (on disk)', 'Ave. SNR',
               'Status', 'Drive', 'Sequence', 'Traj?', 'Pneumo?', 'Blocker / Next action']
    for i, h in enumerate(HEADERS, 1):
        c = st.cell(top, i, h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.border = border
        c.alignment = Alignment(horizontal='center')

    for j, rec in enumerate(recs, 1):
        rr = top + j
        st.cell(rr, 1).value = f'=HYPERLINK("#Sheet1!B{rec["row"]}","row {rec["row"]}")'
        st.cell(rr, 1).font = Font(color='FF0563C1', underline='single')
        st.cell(rr, 2, rec['date'])
        st.cell(rr, 3, rec['sid'])
        disk = st.cell(rr, 4, rec['disk'])
        if rec['disk'] and rec['disk'] != rec['sid']:
            disk.font = Font(bold=True, color='FFC00000')   # ID mismatch
        st.cell(rr, 5, rec['snr'])
        sc = st.cell(rr, 6, rec['status'])
        sc.fill = PatternFill('solid', start_color=FILL[rec['status']])
        st.cell(rr, 7, rec['drive'])
        st.cell(rr, 8, rec['seq'])
        st.cell(rr, 9, rec['traj'])
        st.cell(rr, 10, rec['pneu'])
        st.cell(rr, 11, rec['blocker'])
        for i in range(1, 12):
            st.cell(rr, i).border = border

    # --- notes -----------------------------------------------------------
    nr = top + len(recs) + 2
    st.cell(nr, 1, 'Notes').font = Font(bold=True)
    notes = [
        'Red bold "ID (on disk)" = folder name disagrees with Sheet1. Sheet1 is likely wrong.',
        'NO TRAJECTORY: 10 of these are one sequence, fa_spiral_dyn_fancy_20220812. '
        'Obtaining that single _gp.npy unblocks all 10.',
        'NO DYNAMIC SCAN subjects were never acquired with the spiral-dyn sequence. '
        'No trajectory or code change recovers them — remove from the denominator.',
        'CONTAINER FAIL rows crashed on the OLD image and have not been retried against '
        'the d136eb1 dissolved-phase guard. They may not be true failures.',
        'The d136eb1 image is broken by an unpinned numba upgrade. Pin numpy/scipy/numba '
        'in requirements.txt, rebuild, and bump the sha in workspace/pipeline/recon_codespec.yml.',
    ]
    for i, n in enumerate(notes):
        c = st.cell(nr + 1 + i, 1, '• ' + n)
        c.alignment = Alignment(wrap_text=False)

    # --- geometry --------------------------------------------------------
    widths = [12, 12, 12, 13, 10, 17, 11, 34, 7, 9, 62]
    for i, w in enumerate(widths, 1):
        st.column_dimensions[get_column_letter(i)].width = w
    st.column_dimensions['C'].width = 12
    st.freeze_panes = st.cell(top + 1, 1)
    st.auto_filter.ref = f'A{top}:K{top + len(recs)}'

    wb.save(XLSX)
    print(f'wrote "{name}"  ({len(recs)} ungreened rows, {n_green} green)')
    for s in ORDER:
        print(f'  {s:<16} {counts[s]}')


if __name__ == '__main__':
    main()
